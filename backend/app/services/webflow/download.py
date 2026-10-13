"WPS 个人版取数：**下载整份 xlsx**，不再走剪贴板。"

from __future__ import annotations

import asyncio
import json
import logging
import os
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import httpx

from app.core.config import settings
from app.services.sources.base import SourceError
from app.services.webflow.cdp import Browser, CdpError

log = logging.getLogger(__name__)

DOWNLOADS_ROOT = os.environ.get("WPS_DOWNLOAD_DIR") or str(settings.downloads_dir)

_FILE_INFO_JS = """
(() => {
  try {
    if (window.fileInfo && window.fileInfo.id) {
      return JSON.stringify({
        id: String(window.fileInfo.id),
        name: String(window.fileInfo.name || ''),
        size: Number(window.fileInfo.size || 0),
        modify_time: Number(window.fileInfo.modify_time || 0),
      });
    }
  } catch (e) {}
  try {
    const f = window.__WPSENV__ && window.__WPSENV__.file_info && window.__WPSENV__.file_info.file;
    if (f && f.id) {
      return JSON.stringify({
        id: String(f.id),
        name: String(f.name || ''),
        size: Number(f.size || 0),
        modify_time: Number(f.modify_time || 0),
      });
    }
  } catch (e) {}
  return '';
})()
"""

_FETCH_JS = """
(() => {
  window.__docsyncFetch = null;
  (async () => {
    try {
      const r = await fetch(%s, {credentials: 'include'});
      const text = await r.text();
      window.__docsyncFetch = {status: r.status, text: text.slice(0, 20000)};
    } catch (e) {
      window.__docsyncFetch = {error: String(e).slice(0, 200)};
    }
  })();
  return 'started';
})()
"""


@dataclass
class WpsFile:
    """页面上那个文件（真实 id / 名字 / 字节数 / 最后修改时间）。"""

    file_id: str
    name: str
    size: int
    modify_time: int = 0


async def page_file(browser: Browser, session: str) -> WpsFile:
    """从页面状态里读文件 id。读不到多半是没登录（被踢到登录页）。"""
    try:
        raw = str(await browser.eval(_FILE_INFO_JS, session=session) or "")
    except CdpError:
        raw = ""
    if not raw:
        raise SourceError(
            "网页里读不到文件信息",
            hint="多半是没登录（被踢到登录页了）—— 到应用里重新登录一次；"
            "也可能是金山改版，页面状态里不再放 fileInfo。",
        )
    data = json.loads(raw)
    return WpsFile(
        file_id=str(data.get("id") or ""),
        name=str(data.get("name") or ""),
        size=int(data.get("size") or 0),
        modify_time=int(data.get("modify_time") or 0),
    )


async def _in_page_fetch(browser: Browser, session: str, url: str) -> dict[str, Any]:
    """借页面会话请求一个接口，返回它响应的 JSON。"""
    try:
        await browser.eval(_FETCH_JS % json.dumps(url), session=session)
    except CdpError as exc:
        raise SourceError(f"网页里发请求失败：{exc.message}") from exc
    for _ in range(20):
        await asyncio.sleep(1)
        try:
            got = str(await browser.eval("window.__docsyncFetch", session=session) or "")
        except CdpError:
            continue

        if got and got != "null":
            try:
                payload = json.loads(got) if got.startswith("{") else None
            except json.JSONDecodeError:
                payload = None
            if payload is None:
                raw = str(
                    await browser.eval(
                        "window.__docsyncFetch ? JSON.stringify(window.__docsyncFetch) : ''",
                        session=session,
                    )
                    or ""
                )
                payload = json.loads(raw) if raw else None
            if payload:
                if payload.get("error"):
                    raise SourceError(
                        f"网页里请求 {url} 失败：{payload['error']}",
                        hint="多半是登录态过期，重新登录一次。",
                    )
                if int(payload.get("status") or 0) != 200:
                    raise SourceError(
                        f"下载接口返回 {payload.get('status')}（{url}）",
                        hint="登录态过期，或被踢到登录页。",
                    )
                text = str(payload.get("text") or "{}")
                try:
                    return json.loads(text)
                except json.JSONDecodeError as exc:
                    raise SourceError(f"下载接口没返回 JSON：{text[:200]}") from exc
    raise SourceError("网页里请求超时", hint="网络慢或登录态坏了，重试一次。")


async def download_xlsx(browser: Browser, session: str, app_id: int) -> tuple[Path, WpsFile]:
    """把当前页面的这份文档下载成本地 xlsx，返回 (路径, 文件信息)。"""
    info = await page_file(browser, session)
    payload = await _in_page_fetch(browser, session, f"/api/v3/office/file/{info.file_id}/download")
    url = str(payload.get("download_url") or payload.get("url") or "")
    if not url:
        raise SourceError(
            "下载接口没给文件地址",
            hint="金山改版了：/download 不再返回 download_url，去改映射。",
        )
    dest_dir = Path(DOWNLOADS_ROOT) / f"app-{app_id}"
    dest_dir.mkdir(parents=True, exist_ok=True)
    dest = dest_dir / f"{info.file_id}.xlsx"
    meta = dest.with_name(dest.name + ".meta.json")

    if dest.exists() and dest.stat().st_size > 0:
        same = False
        if meta.exists():
            try:
                old = json.loads(meta.read_text(encoding="utf-8"))
            except (OSError, json.JSONDecodeError):
                old = {}
            version = str(payload.get("fver") or "")
            same_version = bool(version) and version == str(old.get("fver") or "")
            same_timestamp = bool(info.modify_time) and old.get("modify_time") == info.modify_time
            same_size = not info.size or info.size == dest.stat().st_size
            same = same_size and (same_version if version else same_timestamp)
        if same:
            log.info("WPS 下载：文件没变（modify_time=%s），直接用本地那份 %s", info.modify_time, dest)
            return dest, info
    try:
        async with httpx.AsyncClient(timeout=180.0, follow_redirects=True) as client:
            resp = await client.get(url)
            resp.raise_for_status()
            temporary = dest.with_suffix(".xlsx.part")
            temporary.write_bytes(resp.content)
            temporary.replace(dest)

            try:
                meta.write_text(
                    json.dumps(
                        {
                            "file_id": info.file_id,
                            "size": info.size,
                            "modify_time": info.modify_time,
                            "fver": str(payload.get("fver") or ""),
                            "fize": int(payload.get("fize") or 0),
                        }
                    ),
                    encoding="utf-8",
                )
            except OSError:
                pass
    except httpx.HTTPError as exc:
        raise SourceError(
            f"下载文件失败：{exc}",
            hint="签名直链只在有效期内可用；也可能是这台机器连不上金山的存储域名。",
        ) from exc
    log.info(
        "WPS 下载：app-%s 文件 %s（%s 字节）→ %s",
        app_id,
        info.file_id,
        len(dest.read_bytes()),
        dest,
    )
    return dest, info


def _patch_openpyxl_data_validation() -> None:
    "让 openpyxl 读得了「dataValidation 带 id 属性」的 xlsx。"
    from openpyxl.worksheet.datavalidation import DataValidation

    if getattr(DataValidation, "_docsync_id_patched", False):
        return
    original = DataValidation.__init__

    def patched(self, *args, **kwargs):  # noqa: ANN001, ANN002, ANN003
        kwargs.pop("id", None)
        original(self, *args, **kwargs)

    DataValidation.__init__ = patched  # type: ignore[method-assign]
    DataValidation._docsync_id_patched = True  # type: ignore[attr-defined]
    log.info("网页取数：已给 openpyxl 打上 dataValidation.id 兼容补丁")


def _load_workbook(path: Path):
    _patch_openpyxl_data_validation()
    from openpyxl import load_workbook

    return load_workbook(path, read_only=True, data_only=True)


def sheet_grids(path: Path, rows: int = 12) -> dict[str, list[list[str]]]:
    "一次打开文件，把**每张表的前几行**全读出来（列检查、合并都要用）。"
    wb = _load_workbook(path)
    try:
        out: dict[str, list[list[str]]] = {}
        for name in wb.sheetnames:
            ws = wb[name]
            grid: list[list[str]] = []
            for row in ws.iter_rows(values_only=True):
                grid.append(["" if cell is None else str(cell) for cell in row])
                if len(grid) >= rows:
                    break
            out[name] = grid
        return out
    finally:
        wb.close()


def sheet_names(path: Path) -> list[str]:
    """xlsx 里有哪些工作表 —— 这就是**子表清单**（不用再去点"⋯ 全部工作表"）。"""
    wb = _load_workbook(path)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def read_grid(path: Path, sheet: str = "", limit: int | None = None) -> list[list[str]]:
    """按**工作表名**读整块网格（含表头行）。空单元格统一成空串。"""
    wb = _load_workbook(path)
    try:
        if sheet and sheet in wb.sheetnames:
            ws = wb[sheet]
        elif sheet:
            raise SourceError(
                f"这份文档里没有工作表「{sheet}」",
                hint=f"现有工作表：{'、'.join(wb.sheetnames[:20])}",
            )
        else:
            ws = wb[wb.sheetnames[0]]
        grid: list[list[str]] = []
        for row in ws.iter_rows(values_only=True):
            grid.append(["" if cell is None else str(cell) for cell in row])
            if limit is not None and len(grid) >= limit:
                break
        return grid
    finally:
        wb.close()
